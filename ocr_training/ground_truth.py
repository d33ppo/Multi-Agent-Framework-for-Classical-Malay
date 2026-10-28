"""
Ground truth loading for Jawi Qalam Excel files.

Handles the real, observed structure, which is NOT fully consistent across
files (confirmed by inspecting 3 real sample files):
- One row per sentence/paragraph
- Multiple physical printed lines per row, joined by \n within the cell
- Column headers vary: some files name the Jawi ground-truth column
  ("Jawi Correction", "jawi_corrected"), others have an unnamed 4th column
  holding the same kind of data with no header at all.
- Occasional stray/malformed cells: defensively skip empty or non-string
  values rather than crashing.

Given this inconsistency, column identification is done by CONTENT
(which column contains Arabic-script/Jawi text), not by trusting header
names, since header names have already been shown to vary or be missing
across real files in this corpus.
"""
from openpyxl import load_workbook

# Unicode ranges covering Arabic script (used for Jawi) and Arabic
# Presentation Forms, which Jawi extensions sometimes fall into.
ARABIC_RANGES = [
    (0x0600, 0x06FF),  # Arabic
    (0x0750, 0x077F),  # Arabic Supplement
    (0x08A0, 0x08FF),  # Arabic Extended-A
    (0xFB50, 0xFDFF),  # Arabic Presentation Forms-A
    (0xFE70, 0xFEFF),  # Arabic Presentation Forms-B
]


def _is_arabic_char(ch):
    cp = ord(ch)
    return any(lo <= cp <= hi for lo, hi in ARABIC_RANGES)


def _arabic_fraction(text):
    if not text:
        return 0.0
    letters = [c for c in text if c.isalpha()]
    if not letters:
        return 0.0
    arabic_count = sum(1 for c in letters if _is_arabic_char(c))
    return arabic_count / len(letters)


def _find_jawi_column(rows, header, sample_size=8):
    """
    Identifies the Jawi ground-truth column by content, not header name:
    among columns with mostly-Arabic-script text, prefer the LAST such
    column (in every observed file, when there are multiple Jawi-like
    columns such as an auto-transliteration draft plus a human-corrected
    final version, the corrected/final version comes last).

    Column count is taken from data rows, not the header row: at least
    one real file in this corpus has a 4th data column (the actual
    human-corrected text) with no corresponding header cell at all, so
    trusting len(header) would silently miss it and fall back to an
    earlier draft column instead.

    A column only counts as "Jawi-like" if it has real content in
    MULTIPLE sampled rows, not just one. This corpus has been observed
    to contain stray one-off fragments sitting in a phantom extra column
    on a single row (a data-entry artifact); without a minimum-occurrence
    requirement, such a fragment can outrank the real, consistently-used
    ground-truth column just because it happens to be Arabic script.
    """
    sample_rows = rows[1:1 + sample_size]
    n_cols = max([len(header)] + [len(r) for r in sample_rows])

    jawi_like_cols = []
    for col_idx in range(n_cols):
        fractions = []
        for row in sample_rows:
            if col_idx >= len(row):
                continue
            val = row[col_idx]
            if isinstance(val, str) and val.strip():
                fractions.append(_arabic_fraction(val))
        min_occurrences = max(2, len(sample_rows) // 2)
        if len(fractions) >= min_occurrences and (sum(fractions) / len(fractions)) > 0.5:
            jawi_like_cols.append(col_idx)

    if not jawi_like_cols:
        raise ValueError(
            f"No column with predominantly Arabic/Jawi script content found "
            f"consistently across rows. Header was: {header}"
        )

    return jawi_like_cols[-1]


def load_article_lines(xlsx_path):
    """
    Returns a flat list of ground-truth line strings, in reading order,
    for one article (which may span multiple physical PDF pages).
    """
    wb = load_workbook(xlsx_path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header = rows[0]
    jawi_col = _find_jawi_column(rows, header)

    lines = []
    for row in rows[1:]:
        if jawi_col >= len(row):
            continue
        cell = row[jawi_col]
        if not cell or not isinstance(cell, str):
            continue
        for raw_line in cell.split("\n"):
            clean = raw_line.strip()
            if clean:
                lines.append(clean)
    return lines
